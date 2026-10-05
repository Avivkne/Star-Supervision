# -*- coding: utf-8 -*-
"""
report_builder.py
בניית דוח Word מתוך תבנית docxtpl (template.docx).
המודול לא תלוי בטלגרם/OpenAI, כך שאפשר לבדוק אותו לבד:

    python report_builder.py            # בדיקה עצמית עם תבנית זמנית
    python report_builder.py my.docx    # בדיקה עם התבנית האמיתית שלך + נתוני דוגמה
"""
import os
import sys
import tempfile
from datetime import date

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from PIL import Image, ImageOps

TEMPLATE_PATH = os.getenv("TEMPLATE_PATH", "template.docx")
IMAGE_WIDTH_MM = int(os.getenv("IMAGE_WIDTH_MM", "120"))  # רוחב תמונה בדוח
MAX_IMAGES_PER_REMARK = 6
EMPTY = "—"  # מה שמוצג בשדה אופציונלי שלא מולא
SIGNATURE_PATH = os.getenv("SIGNATURE_PATH", "signature.png")  # תמונת חתימה ל-{{ signature_image }}
SIGNATURE_WIDTH_MM = int(os.getenv("SIGNATURE_WIDTH_MM", "35"))

# key -> (תווית בעברית, שדה חובה?)
FIELDS = {
    "visit_date": ("תאריך הביקור", True),
    "project_num": ("מספר פרויקט", True),
    "letter_num": ("מספר מכתב", True),  # מופיע במספר הסימוכין ובתחתית העמוד
    "client_name": ("שם הלקוח", True),
    "contact_person": ("איש קשר", False),
    "client_email": ("אימייל הלקוח", False),
    "structure_name": ("שם המבנה (למשל: ב')", True),
    "inspection_subject": ("נושא הפיקוח (למשל: יציקת תקרה)", True),
    "inspector_name": ("שם המפקח", True),
    "execution_team": ("צוות הביצוע", False),
    "star_present": ("נציגי סטאר מהנדסים שנכחו", False),
    "work_status": ("מצב העבודה", False),
    "author_initials": ("ראשי תיבות באנגלית (למשל A.K)", False),  # אם ריק: נגזר משם המפקח
}


# גיבוי דטרמיניסטי אם המודל לא החזיר ראשי תיבות (התעתיק העיקרי נעשה ב-bot.py על ידי המודל)
_HE2EN = {
    "א": "A", "ב": "B", "ג": "G", "ד": "D", "ה": "H", "ו": "V", "ז": "Z", "ח": "H",
    "ט": "T", "י": "Y", "כ": "K", "ך": "K", "ל": "L", "מ": "M", "ם": "M", "נ": "N",
    "ן": "N", "ס": "S", "ע": "A", "פ": "P", "ף": "P", "צ": "T", "ץ": "T", "ק": "K",
    "ר": "R", "ש": "S", "ת": "T",
}


def derive_initials(name):
    """ראשי תיבות באנגלית בפורמט A.K מתוך שם (עברית או אנגלית): 'Aviv Knebl' -> 'A.K'."""
    letters = []
    for w in (name or "").replace("-", " ").split()[:3]:
        ch = w[0]
        if ch.isascii() and ch.isalpha():
            letters.append(ch.upper())
        elif ch in _HE2EN:
            letters.append(_HE2EN[ch])
    return ".".join(letters)


def with_derived(fields):
    """מחזיר עותק של השדות כשראשי התיבות מושלמים אוטומטית אם לא נאמרו."""
    f = dict(fields)
    if not f.get("author_initials"):
        f["author_initials"] = derive_initials(f.get("inspector_name"))
    return f


def make_composite(paths, out_path):
    """
    ממזג 1..6 תמונות לתמונה אחת (רשת של 2 עמודות) ושומר JPEG.
    כך אפשר לשים כמה תמונות להערה אחת גם כשבתבנית יש רק {{ remark.image }}.
    מחזיר את יחס הגובה/רוחב של התמונה שנוצרה.
    """
    imgs = []
    for p in paths[:MAX_IMAGES_PER_REMARK]:
        with Image.open(p) as im:
            im = ImageOps.exif_transpose(im).convert("RGB")  # מתקן סיבוב מהטלפון
            imgs.append(im.copy())

    if len(imgs) == 1:
        im = imgs[0]
        im.thumbnail((1600, 1600))
        im.save(out_path, "JPEG", quality=85)
        return im.height / im.width

    cell_w, gap = 900, 12
    resized = [im.resize((cell_w, max(1, int(im.height * cell_w / im.width)))) for im in imgs]
    rows = [resized[i:i + 2] for i in range(0, len(resized), 2)]
    row_h = [max(i.height for i in r) for r in rows]
    canvas = Image.new("RGB", (cell_w * 2 + gap, sum(row_h) + gap * (len(rows) - 1)), "white")
    y = 0
    for r, h in zip(rows, row_h):
        x = 0
        for im in r:
            canvas.paste(im, (x, y))
            x += cell_w + gap
        y += h + gap
    canvas.save(out_path, "JPEG", quality=85)
    return canvas.height / canvas.width


def build_report(fields, specific, general, cc, work_dir, template_path=TEMPLATE_PATH,
                 signature_path=SIGNATURE_PATH):
    """
    fields   : dict של שדות (visit_date, project_num, ...)
    specific : list של {"text": str, "images": [נתיבי קבצים]}
    general  : list של str
    cc       : list של str
    מחזיר נתיב ל-docx שנוצר.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"תבנית Word לא נמצאה: {template_path}")

    doc = DocxTemplate(template_path)

    fields = with_derived(fields)
    ctx = {k: (fields.get(k) or EMPTY) for k in FIELDS}
    ctx["author_initials"] = fields.get("author_initials") or ""  # בתחתית העמוד עדיף ריק מ-"—"
    ctx["report_date"] = date.today().strftime("%d/%m/%Y")

    # חתימה: {{ signature_image }}. אם אין קובץ, השדה יוצא ריק ולא שובר את הדוח
    ctx["signature_image"] = ""
    if signature_path and os.path.exists(signature_path):
        ctx["signature_image"] = InlineImage(doc, signature_path, width=Mm(SIGNATURE_WIDTH_MM))

    remarks = []
    for i, r in enumerate(specific, 1):
        image = ""  # מחרוזת ריקה = לא מוצג כלום
        if r.get("images"):
            out = os.path.join(work_dir, f"composite_{i}.jpg")
            aspect = make_composite(r["images"], out)
            # תמונה לאורך (פורטרט) קטנה יותר כדי לא לתפוס עמוד שלם
            width = IMAGE_WIDTH_MM if aspect <= 1.05 else int(IMAGE_WIDTH_MM * 0.65)
            image = InlineImage(doc, out, width=Mm(width))
        remarks.append({"num": i, "text": r.get("text", ""), "image": image})

    ctx["specific_remarks_list"] = remarks
    ctx["general_remarks_list"] = list(general)
    ctx["cc_final_list"] = list(cc)

    # autoescape=True חובה: בלעדיו תו כמו & או < בטקסט שובר את ה-XML של Word
    doc.render(ctx, autoescape=True)

    out_path = os.path.join(work_dir, "report.docx")
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------- בדיקה עצמית
def _sample_data(tmp):
    paths = []
    for n, color in enumerate([(200, 80, 80), (80, 200, 80), (80, 80, 200)], 1):
        p = os.path.join(tmp, f"s{n}.jpg")
        Image.new("RGB", (1200, 900 if n != 3 else 1600), color).save(p)
        paths.append(p)
    fields = {k: f"ערך {k}" for k in FIELDS}
    fields["client_name"] = "חברת א&ב <בע\"מ>"  # בודק escape
    specific = [
        {"text": "נמצא חוסר בכיסוי בטון בקורה K12 בציר 3.", "images": paths[:1]},
        {"text": "יש להשלים זיון בעמוד C5 קומה 2.", "images": paths},  # 3 תמונות
        {"text": "הערה ללא תמונה.", "images": []},
    ]
    return fields, specific, ["התקדמות העבודה תקינה."], ["moshe@example.com"]


def _make_demo_template(path):
    from docx import Document
    d = Document()
    d.add_paragraph("תאריך: {{report_date}} | פרויקט: {{project_num}} | לקוח: {{client_name}}")
    d.add_paragraph("{%p for remark in specific_remarks_list %}")
    d.add_paragraph("{{ remark.num }}. {{ remark.text }}")
    d.add_paragraph("{{ remark.image }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p for note in general_remarks_list %}")
    d.add_paragraph("• {{ note }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p for cc in cc_final_list %}")
    d.add_paragraph("העתק: {{ cc }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("חתימה: {{ signature_image }}  {{ author_initials }}")
    d.save(path)


if __name__ == "__main__":
    tmp = tempfile.mkdtemp(prefix="selftest_")
    tpl = sys.argv[1] if len(sys.argv) > 1 else os.path.join(tmp, "demo_template.docx")
    if len(sys.argv) <= 1:
        _make_demo_template(tpl)
    f, s, g, c = _sample_data(tmp)
    sig = os.path.join(tmp, "signature.png")
    Image.new("RGB", (400, 150), (30, 30, 120)).save(sig)
    result = build_report(f, s, g, c, tmp, template_path=tpl, signature_path=sig)
    print("OK ->", result)
